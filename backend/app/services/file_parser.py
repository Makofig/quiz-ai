"""File parser service: PDF, DOCX, MD, TXT"""

from pathlib import Path
from typing import TextIO
from io import BytesIO
import markdown
from docx import Document
import fitz  # PyMuPDF
import re

SUPPORTED_EXTENSIONS = {"pdf", "docx", "md", "txt"}
MAX_SIZE = 10 * 1024 * 1024  # 10 MB

## Helper function to decode text content with fallback encodings
def decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1")
    
## Helper function to clean up text by removing excessive whitespace and newlines
def clean_text(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()

def parse_pdf(file_content: bytes) -> str:
    doc = fitz.open(stream=file_content, filetype="pdf")
    texts = []
    for page in doc:
        texts.append(page.get_text("text"))
    doc.close()
    return "\n\n".join(texts)


def parse_docx(file_content: bytes) -> str:
    doc = Document(BytesIO(file_content))
    # return "\n\n".join(para.text for para in doc.paragraphs)
    return "\n\n".join(
        para.text.strip()
        for para in doc.paragraphs
        if para.text.strip()
    )


# def parse_md(file_content: bytes | str) -> str:
#     if isinstance(file_content, bytes):
#         file_content = file_content.decode("utf-8")
#     md = markdown.Markdown()
#     return md.convert(file_content)
def parse_md(file_content: bytes | str) -> str:
    return (
        file_content.decode("utf-8")
        if isinstance(file_content, bytes)
        else file_content
    )


def parse_txt(file_content: bytes | str) -> str:
    if isinstance(file_content, bytes):
        return decode_text(file_content)
        # return file_content.decode("utf-8")
    return file_content


def parse_file(file_path: Path) -> str:
    ext = file_path.suffix.lstrip(".").lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported extension: {ext}")
    content = file_path.read_bytes()
    if len(content) > MAX_SIZE:
        raise ValueError(f"File too large ({len(content)} bytes)")
    parsers = {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "md": parse_md,
        "txt": parse_txt,
    }
    return parsers[ext](content)


def parse_file_from_bytes(file_content: bytes, extension: str) -> str:
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported extension: {extension}")
    if len(file_content) > MAX_SIZE:
        raise ValueError(f"File too large ({len(file_content)} bytes)")
    parsers = {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "md": parse_md,
        "txt": parse_txt,
    }
    return clean_text(parsers[extension](file_content))
